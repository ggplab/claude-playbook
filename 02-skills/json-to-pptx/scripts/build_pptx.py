#!/usr/bin/env python3
"""json-to-pptx: 콘텐츠 JSON 한 파일로 업무용 발표 덱(.pptx)을 만든다.

사용:
  python3 build_pptx.py content.json [out.pptx]

원칙:
  - 편집은 PowerPoint에서. 이 스크립트는 "처음 80%"를 빠르게 만든다.
  - 글자 크기는 아래 TYPE 사다리의 이름으로만 고른다. 숫자를 직접 쓰지 않는다.
  - 색은 JSON의 theme 블록이 갖는다. 코드 안에 색을 새로 넣지 않는다.
  - 한 슬라이드에 불릿 6개, 표 8행을 넘으면 경고를 내고 계속 진행한다.

슬라이드 유형 7종: cover, section, bullets, two_column, table, kpi, closing
의존성: python-pptx (pip install python-pptx)
"""
import json
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------- 규격
CANVAS_W, CANVAS_H = 13.333, 7.5          # 16:9 인치
PAD_X, PAD_TOP, PAD_BOTTOM = 0.7, 0.55, 0.45
CONTENT_W = CANVAS_W - PAD_X * 2

# 글자 크기 사다리: (pt, 줄간격). 이름으로만 쓴다.
TYPE = {
    "display": (40, 1.2),   # 표지 제목, 간지 제목
    "title":   (28, 1.25),  # 본문 슬라이드 제목
    "heading": (18, 1.4),   # 열 제목, 표 머리글
    "body":    (16, 1.5),   # 불릿, 표 본문
    "caption": (12, 1.4),   # 보조 설명, 출처
    "chrome":  (10, 1.2),   # 페이지 번호, 푸터
    "kpi":     (34, 1.1),   # KPI 숫자. 4장 카드에 "31,000원"이 한 줄에 들어가는 크기
}

DEFAULT_THEME = {
    "font": "Malgun Gothic",   # Windows 기본 한글 폰트. Mac이면 "Apple SD Gothic Neo"
    "ink": "#111827",          # 본문 글자
    "muted": "#6B7280",        # 보조 글자, 선
    "accent": "#2563EB",       # 강조 한 가지
    "tint": "#EFF6FF",         # 강조의 옅은 면
    "paper": "#FFFFFF",        # 배경
}

LIMITS = {"bullets": 6, "table_rows": 8, "kpi": 4}

_T = dict(DEFAULT_THEME)


def rgb(hexstr):
    h = hexstr.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def warn(msg):
    print(f"  ! {msg}", file=sys.stderr)


# ---------------------------------------------------------------- 기본 도형
def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    return tf


def write(tf, text, tier, *, bold=False, color=None, align=PP_ALIGN.LEFT, first=True):
    size, spacing = TYPE[tier]
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.text = text
    r.font.name = _T["font"]
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = rgb(color or _T["ink"])
    return p


def rect(slide, x, y, w, h, fill, *, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(0.75)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def hline(slide, x, y, w, color=None, weight=0.75):
    ln = slide.shapes.add_connector(1, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = rgb(color or _T["muted"])
    ln.line.width = Pt(weight)
    return ln


def bullets(slide, x, y, w, h, items, *, tier="body", marker="•"):
    tf = textbox(slide, x, y, w, h)
    for i, item in enumerate(items):
        text = item if isinstance(item, str) else item.get("text", "")
        sub = [] if isinstance(item, str) else item.get("sub", [])
        p = write(tf, f"{marker} {text}", tier, first=(i == 0))
        p.space_after = Pt(6)
        for s in sub:
            q = write(tf, f"   – {s}", "caption", color=_T["muted"], first=False)
            q.space_after = Pt(3)
    return tf


def headline(slide, title, subtitle=None):
    """본문 슬라이드 상단: 제목 + 얇은 선. 다음 y를 돌려준다."""
    tf = textbox(slide, PAD_X, PAD_TOP, CONTENT_W, 0.7)
    write(tf, title, "title", bold=True)
    y = PAD_TOP + 0.75
    if subtitle:
        tf2 = textbox(slide, PAD_X, y, CONTENT_W, 0.4)
        write(tf2, subtitle, "caption", color=_T["muted"])
        y += 0.4
    hline(slide, PAD_X, y + 0.05, CONTENT_W)
    return y + 0.35


def chrome(slide, page, total, footer_text):
    y = CANVAS_H - PAD_BOTTOM + 0.05
    if footer_text:
        tf = textbox(slide, PAD_X, y, CONTENT_W - 1.0, 0.3)
        write(tf, footer_text, "chrome", color=_T["muted"])
    tf = textbox(slide, CANVAS_W - PAD_X - 1.0, y, 1.0, 0.3)
    write(tf, f"{page} / {total}", "chrome", color=_T["muted"], align=PP_ALIGN.RIGHT)


def notes(slide, text):
    if text:
        slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- 슬라이드 유형
def slide_cover(prs, d, page, total, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 0.35, CANVAS_H, _T["accent"])
    tf = textbox(s, PAD_X + 0.3, 2.3, CONTENT_W - 0.3, 1.6, anchor=MSO_ANCHOR.BOTTOM)
    write(tf, d.get("title", meta.get("title", "")), "display", bold=True)
    sub = d.get("subtitle", meta.get("subtitle"))
    if sub:
        tf = textbox(s, PAD_X + 0.3, 4.0, CONTENT_W - 0.3, 0.8)
        write(tf, sub, "heading", color=_T["muted"])
    line = " | ".join(v for v in (meta.get("author"), meta.get("date")) if v)
    if line:
        tf = textbox(s, PAD_X + 0.3, 5.6, CONTENT_W - 0.3, 0.5)
        write(tf, line, "caption", color=_T["muted"])
    notes(s, d.get("notes"))
    return s


def slide_section(prs, d, page, total, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, CANVAS_W, CANVAS_H, _T["tint"])
    if d.get("number"):
        tf = textbox(s, PAD_X, 2.2, CONTENT_W, 0.6)
        write(tf, str(d["number"]), "heading", bold=True, color=_T["accent"])
    tf = textbox(s, PAD_X, 2.8, CONTENT_W, 1.4)
    write(tf, d.get("title", ""), "display", bold=True)
    if d.get("subtitle"):
        tf = textbox(s, PAD_X, 4.3, CONTENT_W, 0.8)
        write(tf, d["subtitle"], "heading", color=_T["muted"])
    chrome(s, page, total, meta.get("footer"))
    notes(s, d.get("notes"))
    return s


def slide_bullets(prs, d, page, total, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = headline(s, d.get("title", ""), d.get("subtitle"))
    items = d.get("items", [])
    if len(items) > LIMITS["bullets"]:
        warn(f"p{page} '{d.get('title')}': 불릿 {len(items)}개. {LIMITS['bullets']}개 이하로 나누면 읽힌다")
    bullets(s, PAD_X, y, CONTENT_W, CANVAS_H - y - PAD_BOTTOM - 0.3, items)
    if d.get("source"):
        tf = textbox(s, PAD_X, CANVAS_H - PAD_BOTTOM - 0.35, CONTENT_W, 0.3)
        write(tf, f"출처: {d['source']}", "caption", color=_T["muted"])
    chrome(s, page, total, meta.get("footer"))
    notes(s, d.get("notes"))
    return s


def slide_two_column(prs, d, page, total, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = headline(s, d.get("title", ""), d.get("subtitle"))
    gap = 0.5
    col_w = (CONTENT_W - gap) / 2
    cols = d.get("columns", [])[:2]
    for i, col in enumerate(cols):
        x = PAD_X + i * (col_w + gap)
        accent = bool(col.get("accent"))
        items = col.get("items", [])
        if accent:
            # 강조 면은 내용 높이만큼만. 빈 상자가 바닥까지 내려가지 않게 한다.
            n_sub = sum(len(it.get("sub", [])) for it in items if isinstance(it, dict))
            panel_h = min(0.9 + len(items) * 0.5 + n_sub * 0.35, CANVAS_H - y - PAD_BOTTOM - 0.2)
            rect(s, x - 0.15, y - 0.1, col_w + 0.3, panel_h, _T["tint"], line=_T["accent"])
        tf = textbox(s, x, y, col_w, 0.5)
        write(tf, col.get("heading", ""), "heading", bold=True, color=_T["accent"] if accent else None)
        hline(s, x, y + 0.5, col_w, weight=0.5)
        bullets(s, x, y + 0.7, col_w, CANVAS_H - y - PAD_BOTTOM - 1.0, items)
    chrome(s, page, total, meta.get("footer"))
    notes(s, d.get("notes"))
    return s


def slide_table(prs, d, page, total, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = headline(s, d.get("title", ""), d.get("subtitle"))
    header = d.get("header", [])
    rows = d.get("rows", [])
    if len(rows) > LIMITS["table_rows"]:
        warn(f"p{page} '{d.get('title')}': 표 {len(rows)}행. {LIMITS['table_rows']}행 이하가 읽힌다")
    n_rows, n_cols = len(rows) + 1, max(len(header), max((len(r) for r in rows), default=0))
    row_h = 0.45
    shape = s.shapes.add_table(n_rows, n_cols, Inches(PAD_X), Inches(y), Inches(CONTENT_W), Inches(row_h * n_rows))
    tbl = shape.table
    widths = d.get("widths")  # 비율 리스트, 예: [2, 1, 1]
    if widths and len(widths) == n_cols:
        tot = sum(widths)
        for j, wgt in enumerate(widths):
            tbl.columns[j].width = Inches(CONTENT_W * wgt / tot)

    def fill_cell(cell, text, tier, *, bold=False, color=None, bg=None):
        cell.text = ""
        tf = cell.text_frame
        tf.word_wrap = True
        write(tf, str(text), tier, bold=bold, color=color)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for m in ("margin_left", "margin_right"):
            setattr(cell, m, Inches(0.1))
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(bg or _T["paper"])

    for j in range(n_cols):
        fill_cell(tbl.cell(0, j), header[j] if j < len(header) else "", "heading", bold=True, color=_T["paper"], bg=_T["accent"])
    for i, r in enumerate(rows, start=1):
        for j in range(n_cols):
            val = r[j] if j < len(r) else ""
            fill_cell(tbl.cell(i, j), val, "body", bg=_T["tint"] if i % 2 == 0 else None)
    if d.get("source"):
        tf = textbox(s, PAD_X, CANVAS_H - PAD_BOTTOM - 0.35, CONTENT_W, 0.3)
        write(tf, f"출처: {d['source']}", "caption", color=_T["muted"])
    chrome(s, page, total, meta.get("footer"))
    notes(s, d.get("notes"))
    return s


def slide_kpi(prs, d, page, total, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    y = headline(s, d.get("title", ""), d.get("subtitle"))
    items = d.get("items", [])[:LIMITS["kpi"]]
    if len(d.get("items", [])) > LIMITS["kpi"]:
        warn(f"p{page} '{d.get('title')}': KPI는 {LIMITS['kpi']}개까지만 그린다")
    n = max(len(items), 1)
    gap = 0.4
    w = (CONTENT_W - gap * (n - 1)) / n
    h = 2.6
    for i, it in enumerate(items):
        x = PAD_X + i * (w + gap)
        rect(s, x, y + 0.2, w, h, _T["tint"], line=_T["accent"] if it.get("accent") else None)
        tf = textbox(s, x + 0.3, y + 0.5, w - 0.6, 0.5)
        write(tf, it.get("label", ""), "caption", color=_T["muted"])
        tf = textbox(s, x + 0.3, y + 1.0, w - 0.6, 1.0)
        write(tf, it.get("value", ""), "kpi", bold=True, color=_T["accent"])
        if it.get("delta"):
            tf = textbox(s, x + 0.3, y + 2.1, w - 0.6, 0.5)
            write(tf, it["delta"], "body", color=_T["muted"])
    if d.get("takeaway"):
        tf = textbox(s, PAD_X, y + h + 0.6, CONTENT_W, 1.0)
        write(tf, d["takeaway"], "heading", bold=True)
    chrome(s, page, total, meta.get("footer"))
    notes(s, d.get("notes"))
    return s


def slide_closing(prs, d, page, total, meta):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, 0.35, CANVAS_H, _T["accent"])
    tf = textbox(s, PAD_X + 0.3, 2.2, CONTENT_W - 0.3, 1.2, anchor=MSO_ANCHOR.BOTTOM)
    write(tf, d.get("title", "정리"), "display", bold=True)
    items = d.get("items", [])
    if items:
        bullets(s, PAD_X + 0.3, 3.7, CONTENT_W - 0.3, 2.5, items)
    if d.get("contact"):
        tf = textbox(s, PAD_X + 0.3, CANVAS_H - PAD_BOTTOM - 0.9, CONTENT_W - 0.3, 0.5)
        write(tf, d["contact"], "caption", color=_T["muted"])
    chrome(s, page, total, meta.get("footer"))
    notes(s, d.get("notes"))
    return s


BUILDERS = {
    "cover": slide_cover,
    "section": slide_section,
    "bullets": slide_bullets,
    "two_column": slide_two_column,
    "table": slide_table,
    "kpi": slide_kpi,
    "closing": slide_closing,
}


# ---------------------------------------------------------------- 진입점
def build(content_path, out_path=None):
    content_path = Path(content_path)
    data = json.loads(content_path.read_text(encoding="utf-8"))
    _T.update(data.get("theme", {}))
    slides = data.get("slides", [])
    if not slides:
        sys.exit("slides 배열이 비어 있다")
    total = len(slides)
    meta = {k: data.get(k) for k in ("title", "subtitle", "author", "date", "footer")}

    prs = Presentation()
    prs.slide_width = Inches(CANVAS_W)
    prs.slide_height = Inches(CANVAS_H)

    for i, d in enumerate(slides, start=1):
        kind = d.get("type", "bullets")
        fn = BUILDERS.get(kind)
        if not fn:
            warn(f"p{i}: 모르는 유형 '{kind}'. bullets로 그린다")
            fn = slide_bullets
        fn(prs, d, i, total, meta)

    out = Path(out_path) if out_path else content_path.with_suffix(".pptx")
    prs.save(out)
    print(f"[완료] {out}  ({total}장)")
    return out


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    build(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None)
